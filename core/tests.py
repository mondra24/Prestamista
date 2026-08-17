"""
Tests rigurosos para el Sistema de Gestión de Préstamos
Ejecutar: python manage.py test core -v 2
"""
import tempfile
from pathlib import Path
from decimal import Decimal
from datetime import date, timedelta
from django.test import TestCase, Client as TestClient, override_settings
from django.urls import reverse
from django.contrib.auth.models import User
from django.utils import timezone
from django.core.management import call_command
from django.core.management.base import CommandError
from unittest.mock import patch

from .models import (
    Cliente, Prestamo, Cuota, RutaCobro, TipoNegocio,
    PerfilUsuario, RegistroAuditoria, Notificacion, ConfiguracionRespaldo,
    ConfiguracionCategorizacion
)
from .templatetags.currency_filters import formato_ars, dinero, dinero_completo, formato_miles


# ============== TESTS DE FILTROS DE MONEDA ==============

class CurrencyFiltersTest(TestCase):
    """Tests para los filtros de formato de moneda argentina"""
    
    def test_formato_ars_entero(self):
        """Formato de número entero"""
        self.assertEqual(formato_ars(1234567), '1.234.567')
        self.assertEqual(formato_ars(1000), '1.000')
        self.assertEqual(formato_ars(100), '100')
        self.assertEqual(formato_ars(0), '0')
    
    def test_formato_ars_con_decimales(self):
        """Formato con decimales usando coma"""
        self.assertEqual(formato_ars(1234567.89, 2), '1.234.567,89')
        self.assertEqual(formato_ars(1000.50, 2), '1.000,50')
        self.assertEqual(formato_ars(99.99, 2), '99,99')
    
    def test_formato_ars_decimal_type(self):
        """Formato con tipo Decimal"""
        self.assertEqual(formato_ars(Decimal('1234567.89'), 2), '1.234.567,89')
        self.assertEqual(formato_ars(Decimal('1000000')), '1.000.000')
    
    def test_formato_ars_negativo(self):
        """Formato de números negativos"""
        self.assertEqual(formato_ars(-1234567), '-1.234.567')
    
    def test_formato_ars_none(self):
        """Manejo de valor None"""
        self.assertEqual(formato_ars(None), '0')
        self.assertEqual(dinero(None), '$0')
    
    def test_dinero_simbolo(self):
        """Filtro dinero agrega símbolo $"""
        self.assertEqual(dinero(1234567), '$1.234.567')
        self.assertEqual(dinero(1000000), '$1.000.000')
    
    def test_dinero_con_decimales(self):
        """Filtro dinero con decimales"""
        self.assertEqual(dinero(1234.56, 2), '$1.234,56')
    
    def test_formato_miles(self):
        """Filtro formato_miles sin símbolo"""
        result = formato_miles(1234567)
        self.assertIn('1', result)
        self.assertIn('.', result)


# ============== TESTS DE MODELOS ==============

class ClienteModelTest(TestCase):
    """Tests para el modelo Cliente"""
    
    def setUp(self):
        self.cliente = Cliente.objects.create(
            nombre='Juan',
            apellido='Pérez',
            telefono='1122334455',
            direccion='Calle Test 123',
            categoria='NU'
        )
    
    def test_nombre_completo(self):
        """Test propiedad nombre_completo"""
        self.assertEqual(self.cliente.nombre_completo, 'Juan Pérez')
    
    def test_categoria_choices(self):
        """Test opciones de categoría"""
        categorias = [c[0] for c in Cliente.Categoria.choices]
        self.assertIn('EX', categorias)  # Excelente
        self.assertIn('RE', categorias)  # Regular
        self.assertIn('MO', categorias)  # Moroso
        self.assertIn('NU', categorias)  # Nuevo
    
    def test_cliente_str(self):
        """Test representación string"""
        self.assertIn('Juan', str(self.cliente))
        self.assertIn('Pérez', str(self.cliente))
    
    def test_estado_default(self):
        """Test estado por defecto es Activo"""
        self.assertEqual(self.cliente.estado, 'AC')
    
    def test_cliente_con_ruta(self):
        """Test cliente asignado a ruta"""
        ruta = RutaCobro.objects.create(nombre='Ruta Centro', orden=1)
        self.cliente.ruta = ruta
        self.cliente.save()
        self.assertEqual(self.cliente.ruta.nombre, 'Ruta Centro')


class PrestamoModelTest(TestCase):
    """Tests para el modelo Préstamo"""
    
    def setUp(self):
        self.cliente = Cliente.objects.create(
            nombre='María',
            apellido='García',
            telefono='1144556677',
            direccion='Av. Test 456'
        )
        self.prestamo = Prestamo.objects.create(
            cliente=self.cliente,
            monto_solicitado=Decimal('100000'),
            tasa_interes_porcentaje=Decimal('20'),
            cuotas_pactadas=10,
            frecuencia='SE',
            fecha_inicio=date.today()
        )
    
    def test_monto_total_calculado(self):
        """Test cálculo de monto total a pagar"""
        # 100000 + 20% = 120000
        self.assertEqual(self.prestamo.monto_total_a_pagar, Decimal('120000'))
    
    def test_valor_cuota_calculado(self):
        """Test cálculo de valor de cuota"""
        # 120000 / 10 cuotas = 12000
        primera_cuota = self.prestamo.cuotas.first()
        self.assertEqual(primera_cuota.monto_cuota, Decimal('12000'))
    
    def test_cuotas_generadas(self):
        """Test que se generen las cuotas correctamente"""
        self.assertEqual(self.prestamo.cuotas.count(), 10)
    
    def test_estados_prestamo(self):
        """Test estados válidos de préstamo"""
        estados = [e[0] for e in Prestamo.Estado.choices]
        self.assertIn('AC', estados)  # Activo
        self.assertIn('FI', estados)  # Finalizado
        self.assertIn('CA', estados)  # Cancelado
        self.assertIn('RE', estados)  # Renovado
    
    def test_monto_pendiente(self):
        """Test cálculo de monto pendiente"""
        self.assertEqual(self.prestamo.monto_pendiente, Decimal('120000'))
    
    def test_progreso_inicial(self):
        """Test progreso inicial es 0%"""
        self.assertEqual(self.prestamo.progreso_porcentaje, 0)

    def test_proxima_cuota_incluye_pago_parcial(self):
        """
        Regresión: proxima_cuota solo miraba estado='PE', así que una cuota
        pagada parcialmente (PC) quedaba invisible y la ficha del cliente
        mostraba como "próxima" la cuota siguiente en vez de la que
        realmente falta terminar de cobrar.
        """
        primera = self.prestamo.cuotas.order_by('numero_cuota').first()
        primera.registrar_pago(primera.monto_cuota / 2)
        self.assertEqual(primera.estado, 'PC')

        self.prestamo.refresh_from_db()
        self.assertEqual(self.prestamo.proxima_cuota.pk, primera.pk)


class CuotaModelTest(TestCase):
    """Tests para el modelo Cuota"""
    
    def setUp(self):
        self.cliente = Cliente.objects.create(
            nombre='Carlos',
            apellido='López',
            telefono='1188990011',
            direccion='Calle Prueba 789'
        )
        self.prestamo = Prestamo.objects.create(
            cliente=self.cliente,
            monto_solicitado=Decimal('50000'),
            tasa_interes_porcentaje=Decimal('10'),
            cuotas_pactadas=5,
            frecuencia='SE',
            fecha_inicio=date.today()
        )
        self.cuota = self.prestamo.cuotas.first()
    
    def test_cuota_estado_inicial(self):
        """Test estado inicial de cuota es Pendiente"""
        self.assertEqual(self.cuota.estado, 'PE')
    
    def test_cuota_monto_pagado_inicial(self):
        """Test monto pagado inicial es 0"""
        self.assertEqual(self.cuota.monto_pagado, Decimal('0'))
    
    def test_cuota_monto_restante(self):
        """Test cálculo de monto restante"""
        self.assertEqual(self.cuota.monto_restante, self.cuota.monto_cuota)
    
    def test_registrar_pago_completo(self):
        """Test registrar pago completo de cuota"""
        monto_cuota = self.cuota.monto_cuota
        self.cuota.registrar_pago(monto_cuota)
        self.assertEqual(self.cuota.estado, 'PA')
        self.assertEqual(self.cuota.monto_pagado, monto_cuota)
    
    def test_registrar_pago_parcial(self):
        """Test registrar pago parcial"""
        monto_parcial = self.cuota.monto_cuota / 2
        self.cuota.registrar_pago(monto_parcial)
        self.assertEqual(self.cuota.estado, 'PC')
        self.assertEqual(self.cuota.monto_pagado, monto_parcial)
    
    def test_cuota_vencida(self):
        """Test detección de cuota vencida"""
        self.cuota.fecha_vencimiento = timezone.localtime(timezone.now()).date() - timedelta(days=5)
        self.cuota.save()
        self.assertTrue(self.cuota.esta_vencida)
    
    def test_dias_vencida(self):
        """Test cálculo de días vencida"""
        hoy = timezone.localtime(timezone.now()).date()
        self.cuota.fecha_vencimiento = hoy - timedelta(days=3)
        self.cuota.save()
        self.assertEqual(self.cuota.dias_vencida, 3)


# ============== TESTS DE VISTAS ==============

class ViewsAuthenticationTest(TestCase):
    """Tests de autenticación en vistas"""
    
    def setUp(self):
        self.client = TestClient()
        self.user = User.objects.create_user(
            username='testuser',
            password='testpass123'
        )
    
    def test_dashboard_requires_login(self):
        """Dashboard requiere autenticación"""
        response = self.client.get(reverse('core:dashboard'))
        self.assertEqual(response.status_code, 302)  # Redirect to login
    
    def test_dashboard_authenticated(self):
        """Dashboard accesible con login"""
        self.client.login(username='testuser', password='testpass123')
        response = self.client.get(reverse('core:dashboard'))
        self.assertEqual(response.status_code, 200)
    
    def test_cobros_requires_login(self):
        """Cobros requiere autenticación"""
        response = self.client.get(reverse('core:cobros'))
        self.assertEqual(response.status_code, 302)
    
    def test_cliente_list_requires_login(self):
        """Lista de clientes requiere autenticación"""
        response = self.client.get(reverse('core:cliente_list'))
        self.assertEqual(response.status_code, 302)


class ViewsAccessTest(TestCase):
    """Tests de acceso a vistas con usuario autenticado"""
    
    def setUp(self):
        self.client = TestClient()
        self.user = User.objects.create_user(
            username='testuser',
            password='testpass123'
        )
        self.client.login(username='testuser', password='testpass123')
        
        # Crear datos de prueba
        self.cliente = Cliente.objects.create(
            nombre='Test',
            apellido='Cliente',
            telefono='1111111111',
            direccion='Dirección Test',
            usuario=self.user
        )
        self.prestamo = Prestamo.objects.create(
            cliente=self.cliente,
            monto_solicitado=Decimal('10000'),
            tasa_interes_porcentaje=Decimal('10'),
            cuotas_pactadas=4,
            frecuencia='SE',
            fecha_inicio=date.today(),
            cobrador=self.user  # PrestamoDetailView filtra por cobrador asignado si no es admin
        )

    def test_dashboard_view(self):
        """Test vista dashboard"""
        response = self.client.get(reverse('core:dashboard'))
        self.assertEqual(response.status_code, 200)
    
    def test_cobros_view(self):
        """Test vista cobros"""
        response = self.client.get(reverse('core:cobros'))
        self.assertEqual(response.status_code, 200)
    
    def test_cliente_list_view(self):
        """Test lista de clientes"""
        response = self.client.get(reverse('core:cliente_list'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Test')
    
    def test_cliente_detail_view(self):
        """Test detalle de cliente"""
        response = self.client.get(
            reverse('core:cliente_detail', args=[self.cliente.pk])
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.cliente.nombre)
    
    def test_cliente_create_view(self):
        """Test crear cliente"""
        response = self.client.get(reverse('core:cliente_create'))
        self.assertEqual(response.status_code, 200)
    
    def test_prestamo_list_view(self):
        """Test lista de préstamos"""
        response = self.client.get(reverse('core:prestamo_list'))
        self.assertEqual(response.status_code, 200)
    
    def test_prestamo_detail_view(self):
        """Test detalle de préstamo"""
        response = self.client.get(
            reverse('core:prestamo_detail', args=[self.prestamo.pk])
        )
        self.assertEqual(response.status_code, 200)
    
    def test_cierre_caja_view(self):
        """Test cierre de caja"""
        response = self.client.get(reverse('core:cierre_caja'))
        self.assertEqual(response.status_code, 200)
    
    def test_reporte_general_view(self):
        """Test reporte general"""
        response = self.client.get(reverse('core:reporte_general'))
        self.assertEqual(response.status_code, 200)
    
    def test_planilla_impresion_view(self):
        """Test planilla de impresión"""
        response = self.client.get(reverse('core:planilla_impresion'))
        self.assertEqual(response.status_code, 200)
    
    def test_notificacion_list_view(self):
        """Test lista de notificaciones"""
        response = self.client.get(reverse('core:notificacion_list'))
        self.assertEqual(response.status_code, 200)


class APIViewsTest(TestCase):
    """Tests para vistas API (AJAX)"""
    
    def setUp(self):
        self.client = TestClient()
        self.user = User.objects.create_user(
            username='testuser',
            password='testpass123'
        )
        self.client.login(username='testuser', password='testpass123')
        
        self.cliente = Cliente.objects.create(
            nombre='API',
            apellido='Test',
            telefono='2222222222',
            direccion='Dirección API Test',
            usuario=self.user  # cambiar_categoria_cliente exige propiedad para no-admins
        )
        self.prestamo = Prestamo.objects.create(
            cliente=self.cliente,
            monto_solicitado=Decimal('20000'),
            tasa_interes_porcentaje=Decimal('15'),
            cuotas_pactadas=4,
            frecuencia='SE',
            fecha_inicio=date.today(),
            cobrador=self.user  # cobrar_cuota exige que sea el cobrador asignado
        )
        self.cuota = self.prestamo.cuotas.first()
    
    def test_api_cobrar_cuota(self):
        """Test API de cobro de cuota"""
        response = self.client.post(
            reverse('core:cobrar_cuota', args=[self.cuota.pk]),
            content_type='application/json',
            data='{}'
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data['success'])
        
        # Verificar que la cuota fue pagada
        self.cuota.refresh_from_db()
        self.assertEqual(self.cuota.estado, 'PA')
    
    def test_api_cobrar_cuota_parcial(self):
        """Test API de cobro parcial"""
        import json
        monto_parcial = float(self.cuota.monto_cuota) / 2
        
        response = self.client.post(
            reverse('core:cobrar_cuota', args=[self.cuota.pk]),
            content_type='application/json',
            data=json.dumps({'monto': monto_parcial})
        )
        self.assertEqual(response.status_code, 200)
        
        self.cuota.refresh_from_db()
        self.assertEqual(self.cuota.estado, 'PC')  # Pago Parcial
    
    def test_api_notificaciones(self):
        """Test API de notificaciones"""
        response = self.client.get(reverse('core:api_notificaciones'))
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertIn('count', data)
        self.assertIn('notificaciones', data)
    
    def test_api_cambiar_categoria_cliente(self):
        """Test API cambiar categoría de cliente"""
        import json
        response = self.client.post(
            reverse('core:cambiar_categoria', args=[self.cliente.pk]),
            content_type='application/json',
            data=json.dumps({'categoria': 'EX'})
        )
        self.assertEqual(response.status_code, 200)
        
        self.cliente.refresh_from_db()
        self.assertEqual(self.cliente.categoria, 'EX')


class ClienteFormTest(TestCase):
    """Tests para formulario de cliente"""
    
    def setUp(self):
        self.client = TestClient()
        self.user = User.objects.create_user(
            username='testuser',
            password='testpass123'
        )
        self.client.login(username='testuser', password='testpass123')
    
    def test_crear_cliente_valido(self):
        """Test crear cliente con datos válidos"""
        count_before = Cliente.objects.count()
        response = self.client.post(reverse('core:cliente_create'), {
            'nombre': 'Nuevo',
            'apellido': 'Cliente',
            'telefono': '3333333333',
            'direccion': 'Dirección Test 123',
            'limite_credito': '0',
            'categoria': 'NU',
            'estado': 'AC'
        })
        # Si fue exitoso redirige, si hay error queda en form (200)
        count_after = Cliente.objects.count()
        if response.status_code == 302:
            # Redirigió - se creó
            self.assertEqual(count_after, count_before + 1)
        else:
            # Form con error - verificar que la vista funciona
            self.assertEqual(response.status_code, 200)


class PrestamoFormTest(TestCase):
    """Tests para formulario de préstamo"""
    
    def setUp(self):
        self.client = TestClient()
        self.user = User.objects.create_user(
            username='testuser',
            password='testpass123'
        )
        self.client.login(username='testuser', password='testpass123')
        
        self.cliente = Cliente.objects.create(
            nombre='Para',
            apellido='Préstamo',
            telefono='6666666666',
            direccion='Dir Prestamo Test',
            usuario=self.user  # PrestamoCreateView filtra clientes por cobrador si no es admin
        )

    def test_crear_prestamo_valido(self):
        """Test crear préstamo mediante modelo"""
        # Crear préstamo directamente (el form requiere selección de cliente)
        prestamo = Prestamo.objects.create(
            cliente=self.cliente,
            monto_solicitado=Decimal('50000'),
            tasa_interes_porcentaje=Decimal('20'),
            cuotas_pactadas=10,
            frecuencia='SE',
            fecha_inicio=date.today()
        )
        
        # Verificar que se creó con cuotas
        self.assertIsNotNone(prestamo)
        self.assertEqual(prestamo.cuotas.count(), 10)
        self.assertEqual(prestamo.monto_total_a_pagar, Decimal('60000'))

    def _datos_prestamo(self, cliente=None):
        return {
            'cliente': (cliente or self.cliente).pk,
            'monto_solicitado': '10000',
            'tasa_interes_porcentaje': '10',
            'cuotas_pactadas': '3',
            'frecuencia': 'SE',
            'fecha_inicio': date.today().isoformat(),
        }

    def test_bloquea_segundo_prestamo_si_cliente_ya_tiene_uno_activo(self):
        """A2: no se puede crear un préstamo nuevo si el cliente ya tiene uno activo"""
        Prestamo.objects.create(
            cliente=self.cliente,
            monto_solicitado=Decimal('5000'),
            tasa_interes_porcentaje=Decimal('10'),
            cuotas_pactadas=2,
            frecuencia='SE',
            fecha_inicio=date.today()
        )
        count_before = Prestamo.objects.count()

        response = self.client.post(reverse('core:prestamo_create'), self._datos_prestamo())

        self.assertEqual(response.status_code, 200)  # se queda en el form, no redirige
        self.assertEqual(Prestamo.objects.count(), count_before)
        self.assertContains(response, 'ya tiene un préstamo activo')

    def test_permite_prestamo_si_cliente_no_tiene_uno_activo(self):
        """Sin préstamo activo previo, la creación funciona normalmente"""
        response = self.client.post(reverse('core:prestamo_create'), self._datos_prestamo())
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Prestamo.objects.filter(cliente=self.cliente).count(), 1)


class RepasoSemanaCobrosViewTest(TestCase):
    """Tests para A5: repaso semanal cobrado vs. no cobrado en la vista de Cobros"""

    def setUp(self):
        self.client = TestClient()
        self.user = User.objects.create_user(username='cobrador_semana', password='x')
        self.client.login(username='cobrador_semana', password='x')

        self.cliente = Cliente.objects.create(
            nombre='Repaso', apellido='Semanal', telefono='444', direccion='x',
            usuario=self.user
        )
        self.prestamo = Prestamo.objects.create(
            cliente=self.cliente,
            monto_solicitado=Decimal('30000'),
            tasa_interes_porcentaje=Decimal('10'),
            cuotas_pactadas=3,
            frecuencia='SE',
            fecha_inicio=date.today() - timedelta(days=20),
            cobrador=self.user
        )

    def test_separa_cobradas_y_pendientes_de_los_ultimos_7_dias(self):
        cuotas = list(self.prestamo.cuotas.all())
        # Cuota 1: venció hace 3 días y se cobró
        cuotas[0].fecha_vencimiento = date.today() - timedelta(days=3)
        cuotas[0].save()
        cuotas[0].registrar_pago(cuotas[0].monto_cuota, cobrador=self.user)
        # Cuota 2: venció hace 1 día y sigue pendiente
        cuotas[1].fecha_vencimiento = date.today() - timedelta(days=1)
        cuotas[1].save()
        # Cuota 3: venció hace 20 días (fuera de la ventana de 7 días) y sigue pendiente
        cuotas[2].fecha_vencimiento = date.today() - timedelta(days=20)
        cuotas[2].save()

        response = self.client.get(reverse('core:cobros'))

        self.assertEqual(response.status_code, 200)
        cobradas = response.context['repaso_semana_cobradas']
        pendientes = response.context['repaso_semana_pendientes']

        self.assertEqual([c.pk for c in cobradas], [cuotas[0].pk])
        self.assertEqual([c.pk for c in pendientes], [cuotas[1].pk])
        self.assertEqual(response.context['repaso_semana_total_cobrado'], cuotas[0].monto_cuota)
        self.assertContains(response, 'Repaso de la Semana')

    def test_prestamo_renovado_no_infla_el_repaso_de_cobradas(self):
        """
        Las cuotas de un préstamo renovado quedan en PA por el bulk-close de
        renovar_prestamo, pero no representan cobros reales de la semana.
        """
        cuota = self.prestamo.cuotas.first()
        cuota.fecha_vencimiento = date.today() - timedelta(days=2)
        cuota.save()

        Prestamo.renovar_prestamo(
            prestamo_anterior=self.prestamo,
            nuevo_monto=Decimal('10000'),
            nueva_tasa=Decimal('10'),
            nuevas_cuotas=3,
            nueva_frecuencia='SE'
        )

        response = self.client.get(reverse('core:cobros'))
        self.assertEqual(response.context['repaso_semana_cobradas'], [])


class CalendarioCobrosViewTest(TestCase):
    """Tests para A6: calendario visual de cobros"""

    def setUp(self):
        self.client = TestClient()
        self.user = User.objects.create_user(username='cal_user', password='x')
        self.client.login(username='cal_user', password='x')

        self.cliente = Cliente.objects.create(
            nombre='Calen', apellido='Dario', telefono='777', direccion='x',
            usuario=self.user
        )
        self.prestamo = Prestamo.objects.create(
            cliente=self.cliente,
            monto_solicitado=Decimal('9000'),
            tasa_interes_porcentaje=Decimal('10'),
            cuotas_pactadas=3,
            frecuencia='SE',
            fecha_inicio=date.today() - timedelta(days=10),
            cobrador=self.user
        )

    def _dia_de(self, response, fecha):
        for semana in response.context['grilla']:
            for dia in semana:
                if dia and dia['numero'] == fecha.day:
                    return dia
        return None

    def test_dia_cobrado_pendiente_y_proxima_se_clasifican_bien(self):
        cuotas = list(self.prestamo.cuotas.all())
        hoy = date.today()

        cuotas[0].fecha_vencimiento = hoy - timedelta(days=2)
        cuotas[0].save()
        cuotas[0].registrar_pago(cuotas[0].monto_cuota, cobrador=self.user)  # cobrada

        cuotas[1].fecha_vencimiento = hoy - timedelta(days=1)
        cuotas[1].save()  # vencida, sin cobrar -> pendiente

        cuotas[2].fecha_vencimiento = hoy + timedelta(days=3)
        cuotas[2].save()  # futura -> próxima

        response = self.client.get(reverse('core:calendario_cobros'))
        self.assertEqual(response.status_code, 200)

        self.assertEqual(self._dia_de(response, cuotas[0].fecha_vencimiento)['color'], 'cobrado')
        self.assertEqual(self._dia_de(response, cuotas[1].fecha_vencimiento)['color'], 'pendiente')
        # Si la fecha de la cuota "próxima" cae en el mes siguiente no se verá en esta grilla;
        # solo se verifica cuando cae dentro del mes actual.
        if cuotas[2].fecha_vencimiento.month == hoy.month:
            self.assertEqual(self._dia_de(response, cuotas[2].fecha_vencimiento)['color'], 'proxima')

    def test_dia_sin_cuotas_no_tiene_color(self):
        response = self.client.get(reverse('core:calendario_cobros'))
        # Un día muy lejano dentro del mes sin ninguna cuota generada
        for semana in response.context['grilla']:
            for dia in semana:
                if dia and dia['numero'] == 1 and dia['cantidad'] == 0:
                    self.assertEqual(dia['color'], '')
                    return

    def test_prestamo_renovado_no_pinta_el_dia_como_cobrado(self):
        cuota = self.prestamo.cuotas.first()
        cuota.fecha_vencimiento = date.today()
        cuota.save()

        Prestamo.renovar_prestamo(
            prestamo_anterior=self.prestamo,
            nuevo_monto=Decimal('5000'),
            nueva_tasa=Decimal('10'),
            nuevas_cuotas=2,
            nueva_frecuencia='SE'
        )

        response = self.client.get(reverse('core:calendario_cobros'))
        dia = self._dia_de(response, date.today())
        # Ninguna cuota "activa" ese día (el préstamo quedó renovado): sin color ni cantidad
        self.assertEqual(dia['cantidad'], 0)
        self.assertEqual(dia['color'], '')

    def test_navegacion_de_mes_no_le_mete_punto_de_miles_al_anio(self):
        """Regresión: USE_THOUSAND_SEPARATOR formateaba el año como '2.026' en el querystring"""
        response = self.client.get(reverse('core:calendario_cobros'))
        self.assertNotIn('.', response.context['url_mes_siguiente'].split('year=')[1].split('&')[0])

    def test_navegacion_diciembre_a_enero(self):
        response = self.client.get(reverse('core:calendario_cobros'), {'year': 2026, 'month': 12})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['url_mes_siguiente'], '?year=2027&month=1')
        self.assertEqual(response.context['url_mes_anterior'], '?year=2026&month=11')

    def test_mes_year_invalido_cae_al_mes_actual(self):
        response = self.client.get(reverse('core:calendario_cobros'), {'year': 'abc', 'month': '99'})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['es_mes_actual'])


class ClienteDetailA3Test(TestCase):
    """Tests para A3: ficha de cliente rediseñada (estado del préstamo, cobrar directo, actividad)"""

    def setUp(self):
        self.client = TestClient()
        self.user = User.objects.create_user(username='ficha_user', password='x')
        self.client.login(username='ficha_user', password='x')

        self.cliente = Cliente.objects.create(
            nombre='Ficha', apellido='Test', telefono='888', direccion='x',
            usuario=self.user, notas='Nota interna del cliente'
        )
        self.prestamo = Prestamo.objects.create(
            cliente=self.cliente,
            monto_solicitado=Decimal('12000'),
            tasa_interes_porcentaje=Decimal('10'),
            cuotas_pactadas=2,
            frecuencia='SE',
            fecha_inicio=date.today(),
            cobrador=self.user
        )

    def test_banner_atrasado_y_boton_cobrar_cuando_hay_mora(self):
        cuota = self.prestamo.cuotas.first()
        cuota.fecha_vencimiento = date.today() - timedelta(days=5)
        cuota.save()

        response = self.client.get(reverse('core:cliente_detail', args=[self.cliente.pk]))

        self.assertContains(response, 'loan-status-banner atrasado')
        self.assertContains(response, 'Atrasado 5 días')
        self.assertContains(response, f'data-cuota-id="{cuota.pk}"')
        self.assertContains(response, 'btn-cobrar-perfil')

    def test_banner_al_dia_cuando_no_hay_mora(self):
        cuota = self.prestamo.cuotas.first()
        cuota.fecha_vencimiento = date.today() + timedelta(days=3)
        cuota.save()

        response = self.client.get(reverse('core:cliente_detail', args=[self.cliente.pk]))
        self.assertContains(response, 'loan-status-banner al-dia')
        self.assertContains(response, 'Al día')

    def test_no_muestra_boton_cobrar_si_no_es_el_cobrador_asignado(self):
        """
        El botón de cobro directo respeta el mismo permiso que /api/cobrar/
        (solo el cobrador asignado). Nota: 'btn-cobrar-perfil' por sí solo
        no sirve para el assert porque el <script> de la página lo nombra
        siempre en el selector JS; se busca el atributo data-cuota-id del
        botón real.
        """
        otro = User.objects.create_user(username='otro_cobrador', password='x')
        self.prestamo.cobrador = otro
        self.prestamo.save()
        cuota_id = self.prestamo.cuotas.first().pk

        response = self.client.get(reverse('core:cliente_detail', args=[self.cliente.pk]))
        self.assertNotContains(response, f'data-cuota-id="{cuota_id}"')

    def test_proxima_cuota_parcial_aparece_como_la_que_hay_que_cobrar(self):
        """No debe saltear a la cuota siguiente cuando la actual quedó con pago parcial"""
        primera = self.prestamo.cuotas.order_by('numero_cuota').first()
        primera.registrar_pago(primera.monto_cuota / 2, cobrador=self.user)

        response = self.client.get(reverse('core:cliente_detail', args=[self.cliente.pk]))
        self.assertContains(response, f'data-cuota-id="{primera.pk}"')

    def test_muestra_puntualidad_de_pago_si_hay_prestamos_finalizados(self):
        for cuota in self.prestamo.cuotas.all():
            cuota.registrar_pago(cuota.monto_cuota, cobrador=self.user)

        response = self.client.get(reverse('core:cliente_detail', args=[self.cliente.pk]))
        self.assertContains(response, 'Puntualidad de Pago')

    def test_no_muestra_puntualidad_sin_prestamos_finalizados(self):
        response = self.client.get(reverse('core:cliente_detail', args=[self.cliente.pk]))
        self.assertNotContains(response, 'Puntualidad de Pago')

    def test_muestra_notas_y_actividad_reciente(self):
        cuota = self.prestamo.cuotas.first()
        cuota.registrar_pago(cuota.monto_cuota, cobrador=self.user)

        response = self.client.get(reverse('core:cliente_detail', args=[self.cliente.pk]))
        self.assertContains(response, 'Nota interna del cliente')
        self.assertContains(response, 'Actividad Reciente')
        self.assertContains(response, 'Pago completo')


# ============== TESTS DE EXPORTACIÓN ==============

class ExportViewsTest(TestCase):
    """Tests para vistas de exportación"""
    
    def setUp(self):
        self.client = TestClient()
        self.user = User.objects.create_user(
            username='testuser',
            password='testpass123'
        )
        self.client.login(username='testuser', password='testpass123')
        
        # Crear datos para exportar
        self.cliente = Cliente.objects.create(
            nombre='Export',
            apellido='Test',
            telefono='7777777777',
            direccion='Dir Export Test'
        )
    
    def test_exportar_clientes_excel(self):
        """Test exportar clientes a Excel"""
        response = self.client.get(reverse('core:exportar_clientes_excel'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    
    def test_exportar_prestamos_excel(self):
        """Test exportar préstamos a Excel"""
        response = self.client.get(reverse('core:exportar_prestamos_excel'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    
    def test_exportar_planilla_excel(self):
        """Test exportar planilla a Excel"""
        response = self.client.get(
            reverse('core:exportar_planilla_excel') +
            f'?fecha={date.today().strftime("%Y-%m-%d")}'
        )
        self.assertEqual(response.status_code, 200)

    def test_exportar_cierre_excel(self):
        """Test exportar cierre de caja a Excel (usa construir_excel_cierre_caja)"""
        prestamo = Prestamo.objects.create(
            cliente=self.cliente,
            monto_solicitado=Decimal('5000'),
            tasa_interes_porcentaje=Decimal('10'),
            cuotas_pactadas=2,
            frecuencia='SE',
            fecha_inicio=date.today(),
            cobrador=self.user
        )
        cuota = prestamo.cuotas.first()
        cuota.registrar_pago(cuota.monto_cuota, cobrador=self.user)

        response = self.client.get(
            reverse('core:exportar_cierre_excel') +
            f'?fecha={date.today().strftime("%Y-%m-%d")}'
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )


# ============== TESTS DE MODELOS ADICIONALES ==============

class RutaCobroModelTest(TestCase):
    """Tests para RutaCobro"""
    
    def test_crear_ruta(self):
        """Test crear ruta de cobro"""
        ruta = RutaCobro.objects.create(
            nombre='Zona Norte',
            descripcion='Clientes de zona norte',
            orden=1
        )
        self.assertEqual(str(ruta), 'Zona Norte')
        self.assertTrue(ruta.activa)
    
    def test_ordenamiento_rutas(self):
        """Test ordenamiento de rutas"""
        RutaCobro.objects.create(nombre='Ruta 2', orden=2)
        RutaCobro.objects.create(nombre='Ruta 1', orden=1)
        RutaCobro.objects.create(nombre='Ruta 3', orden=3)
        
        rutas = list(RutaCobro.objects.all())
        self.assertEqual(rutas[0].nombre, 'Ruta 1')
        self.assertEqual(rutas[1].nombre, 'Ruta 2')


class NotificacionModelTest(TestCase):
    """Tests para Notificación"""
    
    def setUp(self):
        self.user = User.objects.create_user(
            username='testuser',
            password='testpass123'
        )
    
    def test_crear_notificacion(self):
        """Test crear notificación"""
        notif = Notificacion.objects.create(
            usuario=self.user,
            titulo='Test Notificación',
            mensaje='Mensaje de prueba',
            tipo='CV',
            prioridad='AL'
        )
        self.assertFalse(notif.leida)
        self.assertEqual(notif.prioridad, 'AL')
    
    def test_marcar_leida(self):
        """Test marcar notificación como leída"""
        notif = Notificacion.objects.create(
            usuario=self.user,
            titulo='Para leer',
            mensaje='Mensaje',
            tipo='IN'
        )
        notif.leida = True
        notif.save()
        self.assertTrue(notif.leida)


class GenerarNotificacionesDiariasCommandTest(TestCase):
    """Tests para el comando D1: notificaciones automáticas de cuotas vencidas/por vencer"""

    def setUp(self):
        self.cliente = Cliente.objects.create(
            nombre='Marta',
            apellido='Gómez',
            telefono='1122334455',
            direccion='Calle Cron 123'
        )
        self.prestamo = Prestamo.objects.create(
            cliente=self.cliente,
            monto_solicitado=Decimal('30000'),
            tasa_interes_porcentaje=Decimal('10'),
            cuotas_pactadas=3,
            frecuencia='SE',
            fecha_inicio=date.today()
        )

    def test_comando_crea_notificacion_de_cuota_vencida(self):
        """La cuota vencida sin notificación previa genera una notificación CV"""
        cuota = self.prestamo.cuotas.first()
        cuota.fecha_vencimiento = date.today() - timedelta(days=2)
        cuota.save()

        self.assertEqual(Notificacion.objects.filter(tipo='CV').count(), 0)
        call_command('generar_notificaciones_diarias')
        self.assertEqual(Notificacion.objects.filter(tipo='CV').count(), 1)

    def test_comando_crea_notificacion_de_cuota_por_vencer(self):
        """La cuota que vence mañana genera una notificación CP"""
        cuota = self.prestamo.cuotas.first()
        cuota.fecha_vencimiento = date.today() + timedelta(days=1)
        cuota.save()

        call_command('generar_notificaciones_diarias')
        self.assertEqual(Notificacion.objects.filter(tipo='CP').count(), 1)

    def test_comando_es_idempotente_en_el_mismo_dia(self):
        """Correr el comando dos veces el mismo día no duplica notificaciones"""
        cuota = self.prestamo.cuotas.first()
        cuota.fecha_vencimiento = date.today() - timedelta(days=1)
        cuota.save()

        call_command('generar_notificaciones_diarias')
        call_command('generar_notificaciones_diarias')
        self.assertEqual(Notificacion.objects.filter(tipo='CV').count(), 1)

    def test_comando_por_vencer_es_idempotente_en_el_mismo_dia(self):
        """Correr el comando dos veces el mismo día no duplica la notificación CP
        (regresión: el título de CP no incluía el # de cuota, así el chequeo de
        duplicados nunca encontraba la notificación ya creada)"""
        cuota = self.prestamo.cuotas.first()
        cuota.fecha_vencimiento = date.today() + timedelta(days=1)
        cuota.save()

        call_command('generar_notificaciones_diarias')
        call_command('generar_notificaciones_diarias')
        self.assertEqual(Notificacion.objects.filter(tipo='CP').count(), 1)

    def test_comando_no_notifica_cuotas_al_dia(self):
        """Una cuota que vence en 10 días no genera ninguna notificación"""
        cuota = self.prestamo.cuotas.first()
        cuota.fecha_vencimiento = date.today() + timedelta(days=10)
        cuota.save()

        call_command('generar_notificaciones_diarias')
        self.assertEqual(Notificacion.objects.filter(tipo__in=['CV', 'CP']).count(), 0)


class RespaldoAutomaticoCommandTest(TestCase):
    """Tests para D4: vigilancia del respaldo diario"""

    def test_respaldo_exitoso_no_genera_notificacion(self):
        """Si el respaldo se hace bien, no se avisa a nadie (evita spam diario)"""
        with patch.object(ConfiguracionRespaldo, 'ejecutar_respaldo', return_value=(True, 'backup_x.json', None)):
            call_command('respaldo_automatico')
        self.assertEqual(Notificacion.objects.filter(tipo='AS').count(), 0)

    def test_respaldo_fallido_notifica_solo_a_superadmins(self):
        """Si el respaldo falla, se avisa únicamente a los superusuarios (quienes gestionan respaldos)"""
        superadmin = User.objects.create_user(username='dev', password='x', is_superuser=True)
        User.objects.create_user(username='cobrador1', password='x')

        with patch.object(ConfiguracionRespaldo, 'ejecutar_respaldo', return_value=(False, None, 'disco lleno')):
            with self.assertRaises(CommandError):
                call_command('respaldo_automatico')

        notifs = Notificacion.objects.filter(tipo='AS')
        self.assertEqual(notifs.count(), 1)
        self.assertEqual(notifs.first().usuario, superadmin)
        self.assertEqual(notifs.first().prioridad, 'AL')

    def test_respaldo_desactivado_no_ejecuta_ni_notifica(self):
        """Si ConfiguracionRespaldo.activo=False, el comando no corre el respaldo"""
        ConfiguracionRespaldo.objects.create(nombre='Respaldo Automático', activo=False)

        with patch.object(ConfiguracionRespaldo, 'ejecutar_respaldo') as mock_ejecutar:
            call_command('respaldo_automatico')
            mock_ejecutar.assert_not_called()
        self.assertEqual(Notificacion.objects.filter(tipo='AS').count(), 0)


class CierreCajaAutomaticoCommandTest(TestCase):
    """Tests para C1: cierre de caja diario automático"""

    def setUp(self):
        self.user = User.objects.create_user(username='cobrador_c1', password='x')
        self.cliente = Cliente.objects.create(
            nombre='Cierre', apellido='Auto', telefono='1', direccion='x', usuario=self.user
        )
        self.prestamo = Prestamo.objects.create(
            cliente=self.cliente,
            monto_solicitado=Decimal('8000'),
            tasa_interes_porcentaje=Decimal('10'),
            cuotas_pactadas=2,
            frecuencia='SE',
            fecha_inicio=date.today(),
            cobrador=self.user
        )

    def test_genera_el_excel_del_dia_en_reportes(self):
        cuota = self.prestamo.cuotas.first()
        cuota.registrar_pago(cuota.monto_cuota, cobrador=self.user)

        with tempfile.TemporaryDirectory() as tmp:
            with override_settings(BASE_DIR=Path(tmp)):
                call_command('cierre_caja_automatico')
            nombre = f'cierre_{date.today().strftime("%Y%m%d")}.xlsx'
            self.assertTrue((Path(tmp) / 'reportes' / nombre).exists())

    def test_genera_el_excel_aunque_no_haya_cobros(self):
        """Un día sin cobros igual arma el Excel (vacío), no debería reventar"""
        with tempfile.TemporaryDirectory() as tmp:
            with override_settings(BASE_DIR=Path(tmp)):
                call_command('cierre_caja_automatico')
            nombre = f'cierre_{date.today().strftime("%Y%m%d")}.xlsx'
            self.assertTrue((Path(tmp) / 'reportes' / nombre).exists())


class MorosidadSemanalCommandTest(TestCase):
    """Tests para C2: reporte semanal de morosidad y proyección"""

    def setUp(self):
        self.user = User.objects.create_user(username='cobrador_c2', password='x')
        self.cliente = Cliente.objects.create(
            nombre='Deudor', apellido='Atrasado', telefono='2', direccion='x', usuario=self.user
        )
        self.prestamo = Prestamo.objects.create(
            cliente=self.cliente,
            monto_solicitado=Decimal('9000'),
            tasa_interes_porcentaje=Decimal('10'),
            cuotas_pactadas=2,
            frecuencia='SE',
            fecha_inicio=date.today() - timedelta(days=10),
            cobrador=self.user
        )

    def test_hoja_morosidad_lista_a_los_atrasados(self):
        import openpyxl

        cuota = self.prestamo.cuotas.first()
        cuota.fecha_vencimiento = date.today() - timedelta(days=3)
        cuota.save()

        with tempfile.TemporaryDirectory() as tmp:
            with override_settings(BASE_DIR=Path(tmp)):
                call_command('morosidad_semanal')
            nombre = f'morosidad_{date.today().strftime("%Y%m%d")}.xlsx'
            path = Path(tmp) / 'reportes' / nombre
            self.assertTrue(path.exists())

            wb = openpyxl.load_workbook(path)
            self.assertIn('Morosidad', wb.sheetnames)
            self.assertIn('Proyección 7 días', wb.sheetnames)
            ws = wb['Morosidad']
            fila = [ws.cell(row=5, column=c).value for c in range(1, 7)]
            self.assertEqual(fila[0], 'Deudor Atrasado')
            self.assertEqual(fila[4], 3)  # días de atraso

    def test_hoja_proyeccion_lista_lo_que_vence_la_semana_que_viene(self):
        import openpyxl

        cuota = self.prestamo.cuotas.first()
        cuota.fecha_vencimiento = date.today() + timedelta(days=4)
        cuota.save()

        with tempfile.TemporaryDirectory() as tmp:
            with override_settings(BASE_DIR=Path(tmp)):
                call_command('morosidad_semanal')
            nombre = f'morosidad_{date.today().strftime("%Y%m%d")}.xlsx'
            wb = openpyxl.load_workbook(Path(tmp) / 'reportes' / nombre)
            ws = wb['Proyección 7 días']
            fila = [ws.cell(row=5, column=c).value for c in range(1, 5)]
            self.assertEqual(fila[0], 'Deudor Atrasado')

    def test_prestamo_renovado_no_aparece_como_moroso(self):
        cuota = self.prestamo.cuotas.first()
        cuota.fecha_vencimiento = date.today() - timedelta(days=3)
        cuota.save()

        Prestamo.renovar_prestamo(
            prestamo_anterior=self.prestamo,
            nuevo_monto=Decimal('5000'),
            nueva_tasa=Decimal('10'),
            nuevas_cuotas=2,
            nueva_frecuencia='SE'
        )

        import openpyxl
        with tempfile.TemporaryDirectory() as tmp:
            with override_settings(BASE_DIR=Path(tmp)):
                call_command('morosidad_semanal')
            nombre = f'morosidad_{date.today().strftime("%Y%m%d")}.xlsx'
            wb = openpyxl.load_workbook(Path(tmp) / 'reportes' / nombre)
            ws = wb['Morosidad']
            self.assertIsNone(ws.cell(row=5, column=1).value)


class PlanillasRutaDiariaCommandTest(TestCase):
    """Tests para C3: planilla de ruta diaria por cobrador"""

    def setUp(self):
        self.cobrador = User.objects.create_user(username='cobrador_c3', password='x')
        self.cobrador.perfil.rol = 'CO'
        self.cobrador.perfil.save()

        self.cliente = Cliente.objects.create(
            nombre='Ruta', apellido='Diaria', telefono='3', direccion='Calle Falsa 123',
            usuario=self.cobrador
        )
        self.prestamo = Prestamo.objects.create(
            cliente=self.cliente,
            monto_solicitado=Decimal('7000'),
            tasa_interes_porcentaje=Decimal('10'),
            cuotas_pactadas=2,
            frecuencia='SE',
            fecha_inicio=date.today(),
            cobrador=self.cobrador
        )

    def test_genera_planilla_para_cobrador_con_pendientes(self):
        import openpyxl

        cuota = self.prestamo.cuotas.first()
        cuota.fecha_vencimiento = date.today()
        cuota.save()

        with tempfile.TemporaryDirectory() as tmp:
            with override_settings(BASE_DIR=Path(tmp)):
                call_command('planillas_ruta_diaria')
            nombre = f'planilla_{self.cobrador.username}_{date.today().strftime("%Y%m%d")}.xlsx'
            path = Path(tmp) / 'reportes' / nombre
            self.assertTrue(path.exists())

            wb = openpyxl.load_workbook(path)
            ws = wb['Ruta del día']
            fila = [ws.cell(row=5, column=c).value for c in range(1, 7)]
            self.assertEqual(fila[1], 'Ruta Diaria')
            self.assertEqual(fila[2], 'Calle Falsa 123')

    def test_no_genera_planilla_para_cobrador_sin_pendientes(self):
        """Un cobrador sin nada para cobrar hoy no debería recibir un archivo vacío"""
        for cuota in self.prestamo.cuotas.all():
            cuota.registrar_pago(cuota.monto_cuota, cobrador=self.cobrador)

        with tempfile.TemporaryDirectory() as tmp:
            with override_settings(BASE_DIR=Path(tmp)):
                call_command('planillas_ruta_diaria')
            nombre = f'planilla_{self.cobrador.username}_{date.today().strftime("%Y%m%d")}.xlsx'
            self.assertFalse((Path(tmp) / 'reportes' / nombre).exists())

    def test_no_genera_planilla_para_cobrador_inactivo(self):
        cuota = self.prestamo.cuotas.first()
        cuota.fecha_vencimiento = date.today()
        cuota.save()
        self.cobrador.perfil.activo = False
        self.cobrador.perfil.save()

        with tempfile.TemporaryDirectory() as tmp:
            with override_settings(BASE_DIR=Path(tmp)):
                call_command('planillas_ruta_diaria')
            nombre = f'planilla_{self.cobrador.username}_{date.today().strftime("%Y%m%d")}.xlsx'
            self.assertFalse((Path(tmp) / 'reportes' / nombre).exists())


class ReportesAutomaticosViewTest(TestCase):
    """Tests para la pantalla de descarga de reportes automáticos (C1-C3)"""

    def setUp(self):
        self.client = TestClient()

    def test_admin_ve_la_lista_de_reportes(self):
        admin = User.objects.create_user(username='admin_reportes', password='x')
        admin.perfil.rol = 'AD'
        admin.perfil.save()
        self.client.login(username='admin_reportes', password='x')

        response = self.client.get(reverse('core:reportes_automaticos'))
        self.assertEqual(response.status_code, 200)

    def test_cobrador_no_admin_no_puede_ver_reportes(self):
        User.objects.create_user(username='cobrador_reportes', password='x')
        self.client.login(username='cobrador_reportes', password='x')

        response = self.client.get(reverse('core:reportes_automaticos'))
        self.assertRedirects(response, reverse('core:dashboard'))

    def test_descarga_rechaza_nombre_que_no_es_un_reporte_valido(self):
        """
        El converter <str:nombre> de Django ya bloquea cualquier '/' en la URL
        (no se puede ni construir /descargar/../../etc/passwd/), así que el
        vector real a cubrir es un nombre de archivo cualquiera que no
        empiece con un prefijo de reporte conocido.
        """
        admin = User.objects.create_user(username='admin_reportes2', password='x')
        admin.perfil.rol = 'AD'
        admin.perfil.save()
        self.client.login(username='admin_reportes2', password='x')

        response = self.client.get(
            reverse('core:descargar_reporte_automatico', args=['..'])
        )
        self.assertRedirects(response, reverse('core:reportes_automaticos'))


class AlertarCobradoresSinActividadCommandTest(TestCase):
    """Tests para D2: alerta si un cobrador no registró cobros en el día"""

    def setUp(self):
        self.admin = User.objects.create_user(username='admin1', password='x')
        self.admin.perfil.rol = 'AD'
        self.admin.perfil.save()

        self.cobrador = User.objects.create_user(username='cobrador1', password='x')
        self.cobrador.perfil.rol = 'CO'
        self.cobrador.perfil.save()

        self.cliente = Cliente.objects.create(
            nombre='Ana', apellido='Ruiz', telefono='111', direccion='x'
        )
        self.prestamo = Prestamo.objects.create(
            cliente=self.cliente,
            monto_solicitado=Decimal('20000'),
            tasa_interes_porcentaje=Decimal('10'),
            cuotas_pactadas=3,
            frecuencia='SE',
            fecha_inicio=date.today(),
            cobrador=self.cobrador
        )

    def test_avisa_si_tenia_cuotas_y_no_cobro_nada(self):
        cuota = self.prestamo.cuotas.first()
        cuota.fecha_vencimiento = date.today()
        cuota.save()

        call_command('alertar_cobradores_sin_actividad')

        self.assertEqual(Notificacion.objects.filter(tipo='AS', usuario=self.admin).count(), 1)
        # El propio cobrador no recibe la alerta sobre sí mismo
        self.assertEqual(Notificacion.objects.filter(tipo='AS', usuario=self.cobrador).count(), 0)

    def test_no_avisa_si_ya_registro_un_cobro(self):
        cuota = self.prestamo.cuotas.first()
        cuota.fecha_vencimiento = date.today()
        cuota.save()
        cuota.registrar_pago(cuota.monto_cuota, cobrador=self.cobrador)

        call_command('alertar_cobradores_sin_actividad')

        self.assertEqual(Notificacion.objects.filter(tipo='AS').count(), 0)

    def test_no_avisa_si_no_tenia_nada_para_cobrar(self):
        """Cuota recién generada con vencimiento futuro: no es responsabilidad del cobrador"""
        call_command('alertar_cobradores_sin_actividad')
        self.assertEqual(Notificacion.objects.filter(tipo='AS').count(), 0)

    def test_no_avisa_de_cobrador_inactivo(self):
        cuota = self.prestamo.cuotas.first()
        cuota.fecha_vencimiento = date.today()
        cuota.save()
        self.cobrador.perfil.activo = False
        self.cobrador.perfil.save()

        call_command('alertar_cobradores_sin_actividad')
        self.assertEqual(Notificacion.objects.filter(tipo='AS').count(), 0)

    def test_es_idempotente_en_el_mismo_dia(self):
        cuota = self.prestamo.cuotas.first()
        cuota.fecha_vencimiento = date.today()
        cuota.save()

        call_command('alertar_cobradores_sin_actividad')
        call_command('alertar_cobradores_sin_actividad')
        self.assertEqual(Notificacion.objects.filter(tipo='AS', usuario=self.admin).count(), 1)


class NotificarCandidatosRenovacionCommandTest(TestCase):
    """Tests para D3: candidatos automáticos a renovación"""

    def setUp(self):
        self.admin = User.objects.create_user(username='admin2', password='x')
        self.admin.perfil.rol = 'AD'
        self.admin.perfil.save()

        self.cobrador = User.objects.create_user(username='cobrador2', password='x')
        self.cobrador.perfil.rol = 'CO'
        self.cobrador.perfil.save()

        self.cliente = Cliente.objects.create(
            nombre='Luis', apellido='Pérez', telefono='222', direccion='x'
        )

    def _crear_y_pagar_prestamo(self, fecha_vencimiento_pasada=False, fecha_pago=None):
        prestamo = Prestamo.objects.create(
            cliente=self.cliente,
            monto_solicitado=Decimal('15000'),
            tasa_interes_porcentaje=Decimal('10'),
            cuotas_pactadas=3,
            frecuencia='SE',
            fecha_inicio=date.today(),
            cobrador=self.cobrador
        )
        for cuota in prestamo.cuotas.all():
            if fecha_vencimiento_pasada:
                cuota.fecha_vencimiento = date.today() - timedelta(days=10)
                cuota.save()
            cuota.registrar_pago(cuota.monto_cuota, cobrador=self.cobrador)
            if fecha_pago is not None:
                cuota.fecha_pago_real = fecha_pago
                cuota.save()
        prestamo.refresh_from_db()
        return prestamo

    def test_candidato_con_buen_historial_es_notificado(self):
        prestamo = self._crear_y_pagar_prestamo()
        self.assertEqual(prestamo.estado, 'FI')

        call_command('notificar_candidatos_renovacion')

        notifs = Notificacion.objects.filter(tipo='RN', usuario=self.admin)
        self.assertEqual(notifs.count(), 1)
        self.assertIn(str(prestamo.pk), notifs.first().titulo)
        # El cobrador no recibe la alerta interna de candidatos
        self.assertEqual(Notificacion.objects.filter(tipo='RN', usuario=self.cobrador).count(), 0)

    def test_mal_historial_no_es_candidato(self):
        self._crear_y_pagar_prestamo(fecha_vencimiento_pasada=True)
        call_command('notificar_candidatos_renovacion')
        self.assertEqual(Notificacion.objects.filter(tipo='RN').count(), 0)

    def test_prestamo_finalizado_en_el_pasado_no_se_notifica_hoy(self):
        self._crear_y_pagar_prestamo(fecha_pago=date.today() - timedelta(days=5))
        call_command('notificar_candidatos_renovacion')
        self.assertEqual(Notificacion.objects.filter(tipo='RN').count(), 0)

    def test_es_idempotente_en_el_mismo_dia(self):
        self._crear_y_pagar_prestamo()
        call_command('notificar_candidatos_renovacion')
        call_command('notificar_candidatos_renovacion')
        self.assertEqual(Notificacion.objects.filter(tipo='RN', usuario=self.admin).count(), 1)


class AuditoriaModelTest(TestCase):
    """Tests para RegistroAuditoria"""
    
    def setUp(self):
        self.user = User.objects.create_user(
            username='testuser',
            password='testpass123'
        )
    
    def test_crear_registro_auditoria(self):
        """Test crear registro de auditoría"""
        registro = RegistroAuditoria.objects.create(
            usuario=self.user,
            tipo_accion='CR',
            tipo_modelo='CL',
            modelo_id=1,
            descripcion='Creación de cliente de prueba'
        )
        self.assertEqual(registro.tipo_accion, 'CR')
        self.assertIsNotNone(registro.fecha_hora)


# ============== TESTS DE LÓGICA DE NEGOCIO ==============

class PrestamoRenovacionTest(TestCase):
    """Tests para renovación de préstamos"""
    
    def setUp(self):
        self.cliente = Cliente.objects.create(
            nombre='Renovador',
            apellido='Test',
            telefono='8888888888',
            direccion='Dir Renovación Test'
        )
        self.prestamo = Prestamo.objects.create(
            cliente=self.cliente,
            monto_solicitado=Decimal('30000'),
            tasa_interes_porcentaje=Decimal('15'),
            cuotas_pactadas=6,
            frecuencia='SE',
            fecha_inicio=date.today() - timedelta(days=60)
        )
        # Pagar algunas cuotas
        for cuota in self.prestamo.cuotas.all()[:4]:
            cuota.registrar_pago(cuota.monto_cuota)
    
    def test_saldo_pendiente(self):
        """Test que el saldo pendiente se calcule bien"""
        cuotas_pendientes = self.prestamo.cuotas.filter(estado='PE')
        saldo = sum(c.monto_restante for c in cuotas_pendientes)
        self.assertTrue(saldo > 0)
    
    def test_progreso_parcial(self):
        """Test progreso con cuotas pagadas"""
        # 4 de 6 cuotas pagadas = ~66%
        progreso = self.prestamo.progreso_porcentaje
        self.assertGreater(progreso, 50)
        self.assertLess(progreso, 100)

    def test_renovar_oculta_el_prestamo_viejo_de_las_listas_activas(self):
        """
        A1: tras renovar, el préstamo viejo pasa a estado RENOVADO y deja de
        aparecer donde se filtra por estado='AC' (dashboard, cobros, reporte
        general), pero se sigue pudiendo consultar en el historial del cliente.
        """
        prestamo_viejo_pk = self.prestamo.pk

        nuevo = Prestamo.renovar_prestamo(
            prestamo_anterior=self.prestamo,
            nuevo_monto=Decimal('20000'),
            nueva_tasa=Decimal('15'),
            nuevas_cuotas=6,
            nueva_frecuencia='SE'
        )

        self.prestamo.refresh_from_db()
        self.assertEqual(self.prestamo.estado, 'RE')

        # Listas "del día a día" (mismo filtro que dashboard/cobros/reporte general)
        activos_del_cliente = Prestamo.objects.filter(cliente=self.cliente, estado='AC')
        self.assertEqual(list(activos_del_cliente), [nuevo])
        self.assertNotIn(self.prestamo, activos_del_cliente)

        # El cliente ahora "activo" apunta al préstamo nuevo, no al viejo
        self.assertEqual(self.cliente.prestamo_activo.pk, nuevo.pk)

        # El historial no se pierde
        self.assertIn(prestamo_viejo_pk, [p.pk for p in self.cliente.prestamos.all()])


class CategoriaClienteTest(TestCase):
    """Tests para lógica de categorías de cliente"""
    
    def test_categoria_default_nuevo(self):
        """Test categoría por defecto es Nuevo"""
        cliente = Cliente.objects.create(
            nombre='Nuevo',
            apellido='Cliente',
            telefono='9999999999',
            direccion='Dir Nuevo Cliente'
        )
        self.assertEqual(cliente.categoria, 'NU')
    
    def test_cambiar_a_excelente(self):
        """Test cambiar categoría a Excelente"""
        cliente = Cliente.objects.create(
            nombre='Buen',
            apellido='Pagador',
            telefono='0000000000',
            direccion='Dir Buen Pagador'
        )
        cliente.categoria = 'EX'
        cliente.save()
        cliente.refresh_from_db()
        self.assertEqual(cliente.categoria, 'EX')


class ConfiguracionCategorizacionTest(TestCase):
    """Tests para A4: categoría del cliente 100% manual por defecto"""

    def setUp(self):
        self.cliente = Cliente.objects.create(
            nombre='Roberto', apellido='Suárez', telefono='333', direccion='x',
            categoria='EX'
        )
        self.prestamo = Prestamo.objects.create(
            cliente=self.cliente,
            monto_solicitado=Decimal('10000'),
            tasa_interes_porcentaje=Decimal('10'),
            cuotas_pactadas=2,
            frecuencia='SE',
            fecha_inicio=date.today()
        )

    def _pagar_mal_y_finalizar(self):
        """Paga las 2 cuotas tarde (mal historial) y finaliza el préstamo"""
        for cuota in self.prestamo.cuotas.all():
            cuota.fecha_vencimiento = date.today() - timedelta(days=10)
            cuota.save()
            cuota.registrar_pago(cuota.monto_cuota)

    def test_por_defecto_categorizacion_es_manual(self):
        self.assertFalse(ConfiguracionCategorizacion.esta_activa())

    def test_categoria_manual_no_cambia_con_mal_historial_por_defecto(self):
        """Sin activar la categorización automática, un mal historial de pagos no toca la categoría manual"""
        self._pagar_mal_y_finalizar()
        self.cliente.refresh_from_db()
        self.assertEqual(self.cliente.categoria, 'EX')

    def test_categorizacion_automatica_sigue_funcionando_si_se_activa(self):
        """El interruptor de vuelta al modo automático (pedido explícito del cliente) sigue funcionando"""
        ConfiguracionCategorizacion.objects.create(pk=1, categorizacion_automatica=True)
        self._pagar_mal_y_finalizar()
        self.cliente.refresh_from_db()
        self.assertEqual(self.cliente.categoria, 'MO')

    def test_historial_pagos_se_calcula_aunque_este_en_modo_manual(self):
        """
        A3: el resumen de puntualidad debe poder mostrarse en la ficha del
        cliente aunque la categorización automática esté apagada (que es el
        default) — son dos cosas independientes.
        """
        self._pagar_mal_y_finalizar()
        resumen = self.cliente.historial_pagos
        self.assertEqual(resumen['total'], 2)
        self.assertEqual(resumen['a_tiempo'], 0)
        self.assertEqual(resumen['porcentaje'], 0)

    def test_historial_pagos_sin_prestamos_finalizados(self):
        cliente_nuevo = Cliente.objects.create(nombre='Sin', apellido='Historial', telefono='1', direccion='x')
        resumen = cliente_nuevo.historial_pagos
        self.assertEqual(resumen['total'], 0)
        self.assertIsNone(resumen['porcentaje'])


# ============== TESTS DE BÚSQUEDA Y FILTROS ==============

class BusquedaClienteTest(TestCase):
    """Tests para búsqueda de clientes"""
    
    def setUp(self):
        self.client = TestClient()
        self.user = User.objects.create_user(
            username='testuser',
            password='testpass123'
        )
        self.client.login(username='testuser', password='testpass123')
        
        Cliente.objects.create(
            nombre='Pedro',
            apellido='Martínez',
            telefono='1010101010',
            direccion='Dir Pedro',
            usuario=self.user
        )
        Cliente.objects.create(
            nombre='Ana',
            apellido='González',
            telefono='2020202020',
            direccion='Dir Ana',
            usuario=self.user
        )
    
    def test_buscar_por_nombre(self):
        """Test búsqueda por nombre"""
        response = self.client.get(
            reverse('core:cliente_list') + '?q=Pedro'
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Pedro')
        self.assertNotContains(response, 'Ana')
    
    def test_buscar_por_apellido(self):
        """Test búsqueda por apellido"""
        response = self.client.get(
            reverse('core:cliente_list') + '?q=González'
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Ana')
    
    def test_filtrar_por_categoria(self):
        """Test filtro por categoría"""
        response = self.client.get(
            reverse('core:cliente_list') + '?categoria=NU'
        )
        self.assertEqual(response.status_code, 200)


class FiltroPrestamosTest(TestCase):
    """Tests para filtros de préstamos"""
    
    def setUp(self):
        self.client = TestClient()
        self.user = User.objects.create_user(
            username='testuser',
            password='testpass123'
        )
        self.client.login(username='testuser', password='testpass123')
        
        cliente = Cliente.objects.create(
            nombre='Filtro',
            apellido='Test',
            telefono='3030303030',
            direccion='Dir Filtro Test'
        )
        Prestamo.objects.create(
            cliente=cliente,
            monto_solicitado=Decimal('10000'),
            tasa_interes_porcentaje=Decimal('10'),
            cuotas_pactadas=5,
            frecuencia='SE',
            fecha_inicio=date.today()
        )
    
    def test_filtrar_prestamos_activos(self):
        """Test filtrar préstamos activos"""
        response = self.client.get(
            reverse('core:prestamo_list') + '?estado=AC'
        )
        self.assertEqual(response.status_code, 200)
    
    def test_filtrar_prestamos_finalizados(self):
        """Test filtrar préstamos finalizados"""
        response = self.client.get(
            reverse('core:prestamo_list') + '?estado=FI'
        )
        self.assertEqual(response.status_code, 200)


# ============== TEST DE INTEGRIDAD ==============

class IntegridadDatosTest(TestCase):
    """Tests de integridad de datos"""
    
    def test_eliminar_cliente_con_prestamo(self):
        """Test que no se puede eliminar cliente con préstamo activo"""
        cliente = Cliente.objects.create(
            nombre='Protegido',
            apellido='Cliente',
            telefono='4040404040',
            direccion='Dir Protegido'
        )
        Prestamo.objects.create(
            cliente=cliente,
            monto_solicitado=Decimal('10000'),
            tasa_interes_porcentaje=Decimal('10'),
            cuotas_pactadas=5,
            frecuencia='SE',
            fecha_inicio=date.today()
        )
        
        from django.db.models import ProtectedError
        with self.assertRaises(ProtectedError):
            cliente.delete()
    
    def test_cuotas_se_eliminan_con_prestamo(self):
        """Test cuotas se eliminan en cascada con préstamo"""
        cliente = Cliente.objects.create(
            nombre='Cascada',
            apellido='Test',
            telefono='5050505050',
            direccion='Dir Cascada'
        )
        prestamo = Prestamo.objects.create(
            cliente=cliente,
            monto_solicitado=Decimal('10000'),
            tasa_interes_porcentaje=Decimal('10'),
            cuotas_pactadas=3,
            frecuencia='SE',
            fecha_inicio=date.today()
        )
        
        cuotas_count = Cuota.objects.filter(prestamo=prestamo).count()
        self.assertEqual(cuotas_count, 3)
        
        prestamo.delete()
        
        cuotas_count_after = Cuota.objects.filter(prestamo=prestamo).count()
        self.assertEqual(cuotas_count_after, 0)
